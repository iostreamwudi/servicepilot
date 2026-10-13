from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "docs" / "ServicePilot_PRD_v1.0.docx"

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(89, 99, 110)
LIGHT = "F2F4F7"
CALLOUT = "E8EEF5"
WHITE = RGBColor(255, 255, 255)
PAGE_WIDTH_DXA = 9360

# Arial Unicode MS is installed locally and is visible to both Word and the
# isolated LibreOffice renderer, so Chinese glyphs survive deterministic QA.
CN_FONT = "Arial Unicode MS"

def set_font(run, size=11, bold=None, color=None, name=CN_FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_cell_width(cell, width):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(width))
    tcW.set(qn("w:type"), "dxa")

def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120")
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    trPr.append(cant_split)

def set_paragraph(p, before=0, after=6, line=1.10, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align

def add_para(doc, text="", style=None, before=0, after=6, bold_prefix=None, color=None, size=None):
    p = doc.add_paragraph(style=style)
    set_paragraph(p, before, after)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_font(first, size or 11, True, color)
        rest = p.add_run(text[len(bold_prefix):])
        set_font(rest, size or 11, False, color)
    else:
        run = p.add_run(text)
        set_font(run, size or 11, None, color)
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    set_paragraph(p, 0, 5, 1.12)
    set_font(p.add_run(text), 10.5)
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    set_paragraph(p, 0, 5, 1.12)
    set_font(p.add_run(text), 10.5)
    return p

def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p

def write_cell(cell, text, bold=False, color=None, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph(p, 0, 0, 1.05)
    run = p.add_run(text)
    set_font(run, size, bold, color)

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for cell, header in zip(hdr.cells, headers):
        shade(cell, LIGHT)
        write_cell(cell, header, bold=True, color=INK, size=9)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for cell, value in zip(cells, row_values):
            write_cell(cell, value, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table

def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=2)
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    table.style = "Table Grid"
    set_table_geometry(table, [1320, 8040])
    shade(table.cell(0, 0), "1F4D78")
    shade(table.cell(0, 1), CALLOUT)
    write_cell(table.cell(0, 0), label, bold=True, color=WHITE, size=9)
    write_cell(table.cell(0, 1), text, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_metadata(doc):
    rows = [
        ("文档状态", "V3.2 - Phase 1.1 Real-model Retrieval Eval"),
        ("产品", "ServicePilot 企业智能服务运营平台"),
        ("文档负责人", "产品负责人（待指定）"),
        ("适用范围", "青屿商城生产空间 + ServicePilot 沙箱"),
        ("最后更新", "2026-08-25"),
        ("维护规则", "每次功能范围、角色权限、指标或发布逻辑变更后同步更新本 PRD。"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        set_paragraph(p, 0, 2, 1.05)
        set_font(p.add_run(label + "："), 10, True, INK)
        set_font(p.add_run(value), 10, False, MUTED)

def setup(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = CN_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK, 8, 4),
    ]:
        s = styles[style_name]
        s.font.name = CN_FONT
        s._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
        s.font.size = Pt(size)
        s.font.color.rgb = color
        s.font.bold = True
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph(header, 0, 0, 1.0)
    set_font(header.add_run("SERVICEPILOT | 产品需求文档"), 8.5, True, MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph(footer, 0, 0, 1.0)
    set_font(footer.add_run("内部工作稿 | 2026-08-25"), 8.5, False, MUTED)

def main():
    doc = Document()
    setup(doc)
    # Cover / masthead
    p = doc.add_paragraph()
    set_paragraph(p, 18, 4, 1.0)
    set_font(p.add_run("PRODUCT REQUIREMENTS DOCUMENT"), 10, True, BLUE)
    p = doc.add_paragraph()
    set_paragraph(p, 0, 4, 1.0)
    set_font(p.add_run("ServicePilot 企业智能服务运营平台"), 24, True, INK)
    p = doc.add_paragraph()
    set_paragraph(p, 0, 16, 1.1)
    set_font(p.add_run("角色工作台、可信知识治理与 Agent 服务闭环"), 14, False, MUTED)
    add_metadata(doc)
    add_callout(doc, "定位", "面向企业内部服务团队的 AI 服务运营平台：让每一次 Agent 回答都具备可信依据、明确边界、可追溯反馈和可控发布。")
    add_heading(doc, "阅读说明", 2)
    add_para(doc, "本 PRD 以当前已部署原型为基线，既记录现有能力，也定义下一阶段可上线产品需要满足的要求。界面中的演示数值仅用于说明信息架构与评测方式，不应被视为真实经营数据。", after=8)
    add_para(doc, "[Assumption] 当前尚未完成正式用户研究与真实生产埋点校验；所有目标值在试点前须由业务、数据和合规负责人确认。", after=4, color=MUTED)
    doc.add_page_break()

    add_heading(doc, "1. 执行摘要", 1)
    add_para(doc, "ServicePilot 用于连接消费者咨询、企业知识、AI Agent、人工客服与治理流程。它解决的不是单一聊天机器人问题，而是企业在规模化使用生成式 AI 时常见的三类断裂：回答缺少依据、人工接管缺少上下文、优化没有闭环。产品以角色工作台组织工作：一线客服完成服务交付，主管控制风险与 SLA，知识运营维护可信内容，AI 运营改进模型与策略，系统管理员控制权限和审计。")
    add_callout(doc, "本期决策", "先以“问答 - 知识 - 评测 - 发布”闭环为核心，优先完善知识运营与 AI 运营；消费者端仅作为真实问答输入和服务结果验证，不作为后台主用户。")

    add_heading(doc, "2. 背景与问题定义", 1)
    add_heading(doc, "2.1 目标用户的问题", 2)
    for item in [
        "一线客服：接管 AI 失败或高风险会话时，往往需要反复查订单、规则和历史，用户被迫重复描述。",
        "服务主管：难以及时发现 SLA 风险、知识错误和 AI 失败之间的关联，调度依赖人工经验。",
        "知识运营：业务资料分散、版本冲突、责任不清，发布后无法证明回答是否真的改善。",
        "AI 运营：无法快速判断失败来自意图识别、检索、知识、策略还是模型回答，优化容易只看单点指标。",
        "企业治理方：需要控制谁能修改策略、查看敏感数据、发布知识，以及在出错时追溯完整链路。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "2.2 当前原型证据与限制", 2)
    add_para(doc, "原型已实现角色登录选择、会话与工单协作、已发布知识引用、知识运营的测试/复核/灰度流程、质检评测、策略配置、权限与审计页面。本次增量将体验用户端的提问、Agent 回复、知识引用、会话状态和转人工任务写入持久化数据，并让客服工作台从数据库读取和更新会话。工作台指标和部分审批状态仍不代表完整生产数据。")
    add_para(doc, "[Open Question] 需要在试点前补充访谈、日志和业务指标，验证各角色的实际高频任务、决策权限与目标阈值。", color=MUTED)

    add_heading(doc, "3. 产品目标、范围与非目标", 1)
    add_heading(doc, "3.1 产品目标", 2)
    for item in [
        "让 Agent 的回答可追溯到已批准知识，并能识别不可自动处理的边界。",
        "让人工客服在接管时获得完整上下文、可核验事实与可编辑建议。",
        "让知识和策略变更必须经过测试、审批和可回退的发布流程。",
        "让运营人员从失败案例回溯到知识、策略或服务流程，并形成可量化改进闭环。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "3.2 本期范围", 2)
    for item in [
        "企业登录、工作空间隔离、授权工作台选择与角色切换。",
        "一线客服会话接管、上下文核验、AI 协助、工单协同、结案与知识引用。",
        "服务主管运营总览、风险优先级、SLA 关注与跨角色协同入口。",
        "知识运营工作队列、资产库、来源导入、冲突识别、回答验证、双人复核与灰度发布原型。",
        "AI 运营质检评测、失败复核、策略中心、效果与成本分析。",
        "系统管理员资源权限、安全审计和策略治理界面。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "3.3 非目标", 2)
    for item in [
        "不在本期直接替代订单、支付、物流、CRM 或工单底层系统。",
        "不允许 Agent 无审批执行退款、赔付、账户或地址等高风险业务动作。",
        "不在本期建设面向消费者的独立后台；消费者端仅承担问答输入和体验验证。",
        "不将演示用测试、复核、灰度按钮描述为已持久化的企业审批系统。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "4. 用户、角色与权限模型", 1)
    add_para(doc, "消费者是服务对象，不是本产品后台主用户。其咨询、凭证、满意度和追问构成改进闭环的输入。以下内部角色根据登录时的授权进入不同工作台。")
    add_table(doc,
        ["角色", "核心目标", "当前工作台与操作", "边界"],
        [
            ["一线客服", "安全、高效地解决需要人工介入的问题", "会话接管；查看订单/用户/AI 摘要；采用或编辑建议回复；建单、升级、结案；查询并引用生效知识", "不可编辑知识、发布策略或绕过审批执行高风险业务动作"],
            ["服务主管", "控制服务质量、SLA 与高风险事项", "运营总览；风险排序；工单分配/催办；查看服务趋势与跨角色协作记录", "不直接替代知识或模型负责人进行内容/策略发布"],
            ["知识运营", "把可信资料变为可验证的 Agent 知识", "缺口/冲突/复核队列；资产库；来源导入；内容边界；真实问法测试；双人复核；灰度发布", "不决定模型路由；业务与合规审批不能被单人绕过"],
            ["AI 运营", "提高识别、检索、回答和路由质量", "质检评测；失败复核；知识关联；策略配置；质量、成本与时延分析", "不替代业务方定义规则，也不直接处理用户会话"],
            ["系统管理员", "保障权限、审计与生产变更可控", "资源权限；安全审计；策略治理；权限冲突与到期复核", "不拥有业务规则内容的最终解释权"],
        ], [1200, 1800, 4200, 2160])
    add_para(doc, "[Open Question] 业务规则负责人和合规审批人已在双人复核流程中出现，但尚未作为独立登录工作台实现；正式上线前需要明确其权限和审批责任。", color=MUTED)

    add_heading(doc, "5. 端到端业务闭环", 1)
    add_number(doc, "消费者发起咨询，Agent 识别意图、检索可信知识、执行风险校验并生成回答。")
    add_number(doc, "低置信度、高风险、需要身份核验或无法承诺的场景转交一线客服，并保留会话摘要、事实和知识版本。")
    add_number(doc, "一线客服处理、建单、升级或结案；其对建议、摘要和知识的反馈进入质检与治理队列。")
    add_number(doc, "服务主管按 SLA、风险和业务影响调度资源，并推动跨角色问题处理。")
    add_number(doc, "知识运营修复缺口与冲突；AI 运营定位评测失败并调整检索、模型或策略。")
    add_number(doc, "经过测试、双人复核和灰度观察后发布；管理员对权限和审计链路进行治理。")
    add_callout(doc, "闭环原则", "任何知识或策略变更都必须可定位来源、可验证回答、可审批发布、可观察影响、可回退。")

    add_heading(doc, "6. 解决方案与功能需求", 1)
    add_heading(doc, "6.1 登录、空间与角色工作台", 2)
    add_table(doc, ["编号", "需求", "优先级", "验收标准"], [
        ["FR-01", "用户登录后选择企业空间，再选择已授权工作台；可记住选择并从账号菜单切换。", "P0", "仅显示该空间的授权角色；切换角色后首页、导航和数据入口随角色变化；无授权角色不可进入。"],
        ["FR-02", "每个角色首页展示其目标、优先任务、关键指标和高频入口。", "P0", "一线客服、主管、知识运营、AI 运营、管理员看到不同待办与导航；任务按影响/风险/时限组织。"],
        ["FR-03", "企业空间之间数据和授权隔离。", "P0", "切换空间后，当前角色必须重新校验授权；界面明确展示空间与环境。"],
        ["FR-04", "消费者端真实提问创建可追溯会话，并保存用户消息、Agent 回复、知识引用和处理轨迹。", "P0", "每次请求写入会话与消息；响应返回会话编号；不会只停留在前端状态。"],
        ["FR-05", "Agent 对低置信度、高风险或规则要求的请求生成转人工任务。", "P0", "会话状态变为待接管；客服刷新或轮询后能看到用户原问、AI 回复、依据和摘要。"],
        ["FR-06", "消费者端以真实客服会话呈现多角色测试入口，支持长期会话列表、历史消息恢复和人工回复回传。", "P0", "用户可选择测试身份、切换历史会话或新建会话；用户端不展示 Agent/模型内部术语；客服消息在同一会话中实时可见；历史记录不因短期状态过期而消失。"],
        ["FR-07", "生产租户与测试租户隔离；生产租户默认从空数据开始，测试种子仅通过显式测试通道加载。", "P0", "生产请求不会读取演示知识、会话、工单或业务事实；响应和数据库记录带租户与来源标识。"],
        ["FR-08", "每次 Agent 请求记录 traceId、模型、知识来源、工具来源、失败码和数据新鲜度。", "P0", "可从用户消息、客服消息、工单和审计记录追溯同一请求；不记录 API Key；合成数据不得标记为生产真实数据。"],
    ], [850, 3700, 700, 4110])
    add_heading(doc, "6.2 一线客服工作台", 2)
    add_table(doc, ["编号", "需求", "优先级", "验收标准"], [
        ["FR-10", "以待接管、处理中、等待内部、已解决组织会话队列。", "P0", "客服可筛选状态并查看剩余 SLA、风险、情绪、来源渠道和交接原因。"],
        ["FR-11", "会话详情展示客户消息、AI 摘要、已核验事实、订单/用户上下文和脱敏提示。", "P0", "事实必须标识来源；客服可反馈摘要不准确；敏感字段不得完整暴露。"],
        ["FR-12", "AI 生成带知识依据的推荐回复，并支持人工采用、编辑或反馈无帮助。", "P0", "建议显示知识依据；采用后进入回复框；高风险场景明确提示不能自动承诺。"],
        ["FR-13", "客服可创建/跟进工单、升级主管并结构化结案。", "P0", "结案要求选择结果，确认用户已解决后才能关闭；结案记录可进入 SLA 与质检。"],
        ["FR-14", "客服只可查询和引用已发布知识。", "P0", "搜索结果仅返回 active 知识；引用时记录知识编号与版本；无编辑、审核、发布入口。"],
        ["FR-15", "客服接管、发送人工回复、结案和重新打开均持久化，并写入审计事件。", "P0", "刷新页面后状态与消息仍保留；接管和结案不依赖前端内存。"],
        ["FR-16", "客服回复写入消息表并回传至消费者端；数据库写入失败时接口返回失败。", "P0", "用户端轮询或 SSE 能收到客服消息；会话历史按时间完整展示，不被最新一条消息覆盖。"],
    ], [850, 3700, 700, 4110])
    add_heading(doc, "6.3 知识运营工作台", 2)
    add_table(doc, ["编号", "需求", "优先级", "验收标准"], [
        ["FR-20", "按受影响会话、风险和到期时间排序知识缺口、口径冲突、待复核和即将过期任务。", "P0", "工作队列可定位关联知识与任务来源；默认优先处理高影响高风险项。"],
        ["FR-21", "知识资产包含客服可引用内容、场景、责任团队、风险、接管策略、来源、版本和状态。", "P0", "编辑仅保存草稿，不影响线上；来源可追溯；冲突被显式标记。"],
        ["FR-22", "导入官方政策、内部规范或 FAQ，系统生成候选知识和冲突提示。", "P1", "导入不能绕过人工审核直接发布；显示来源身份、预计片段数和潜在冲突。"],
        ["FR-23", "使用真实用户问法完成召回、边界和转人工验证。", "P0", "测试通过前“提交双人复核”不可用；结果展示命中知识、边界与接管检查。"],
        ["FR-24", "通过业务与合规双人复核后进行灰度发布，并可回退线上版本。", "P0", "发布状态区分草稿/复核中/灰度中；展示流量比例与回退能力。"],
        ["FR-25", "维护版本化的预设回复/话术资产，包含客户真实问法、推荐回复、避免表达、回复结构、适用场景和责任人。", "P0", "知识运营可编辑并提交复核；客服可反馈但不可发布；草稿不影响线上回答；内容可追溯到版本。"],
        ["FR-26", "将已发布话术作为 Agent 的风格与表达上下文，并提供训练影响、风险词扫描、双人审批和回滚状态。", "P0", "发布后仅影响表达方式，不覆盖业务知识与安全策略；主管可批准/发布/回滚；AI 运营可观察覆盖会话与效果。"],
    ], [850, 3700, 700, 4110])
    add_heading(doc, "6.4 AI 运营与策略中心", 2)
    add_table(doc, ["编号", "需求", "优先级", "验收标准"], [
        ["FR-30", "提供全量评测、回归、盲测、新数据泛化和人工复核队列。", "P0", "已知回归与新数据泛化必须使用独立版本和指标展示；结果能按识别、检索、回答、策略等维度定位失败；人工结论保留原始样本和复核信息。"],
        ["FR-31", "失败案例可回流知识、策略或数据治理工作项。", "P0", "运营人员能明确标记失败归因；不允许直接覆盖原始标注，审批后形成新版本。"],
        ["FR-32", "策略中心配置低置信度兜底、高风险转人工、上下文继承和模型路由。", "P0", "生产策略发布前显示检查项和双人审批状态；高风险策略默认不允许自动承诺。"],
        ["FR-33", "数据分析联动服务、质量、成本和时延指标。", "P1", "支持会话解决漏斗和质量 x 成本观察；数据口径需可追溯。"],
        ["FR-34", "在不切换生产答案来源的前提下，后台运行版本化 Hybrid 影子检索并记录候选差异。", "P0", "按租户、Embedding 模型和当前索引版本读取；记录现网/影子来源、Top-K 分数与差值、选择理由、索引版本、耗时和失败码；影子失败不改变回答、Prompt 或转人工。"],
        ["FR-35", "授权质量角色复核真实影子候选变化，并以同一索引版本的分层证据判断是否可进入灰度设计。", "P0", "仅 completed 且 changed 的样本可复核；保留改善/退化/等价/无法判断、正确来源、复核人、时间和审计；按类目/版本统计，门禁只读且不得自动切流。"],
        ["FR-36", "受约束 Hybrid 作为完整 Agent 技术基线，并保留稳定来源回退。", "P0", "候选按租户、Embedding模型、当前索引版本、active状态和业务类别隔离；向量或候选读取失败时使用稳定来源；新数据完整Agent为主指标，旧集用于明显退化和高风险保护；技术基线不等于自动切生产。"],
        ["FR-37", "检索证据不足时只允许一次有界Corrective RAG，不允许自由循环。", "P0", "最多执行一次结构化查询改写与重检索；改写新增词不得充当采用证据；仍受租户、模型、索引版本、active状态和类目约束；第二次仍不足时回退83%技术基线并记录次数。"],
        ["FR-38", "单条知识不足时只组合少量互补知识证据。", "P0", "证据必须来自同租户、同索引版本、active状态和同业务类目；组合前检查适用条件、版本、来源与冲突；冲突或证据不足时不组合并回退单知识基线。"],
        ["FR-39", "实时事实、知识规则和模型声明通过有界Agent统一治理。", "P0", "实时状态仅可来自授权只读工具；缺少订单号先追问；知识与工具证据保留类型、来源、版本/新鲜度；生成后审查无证据数字、承诺、状态冲突与未授权执行；单次最多1次工具、1次模型、7个步骤和25秒，超限安全降级。"],
        ["FR-47", "Step 9.8 V2 在显式实验开关下使用 RetrievalScope、BM25/Dense 宽召回、RRF、业务动作 Boost 和语义 Reranker 形成最终证据候选。", "P0", "Router 只能提供检索范围与偏好，不能锁定唯一知识；V1 和普通生产行为不变；V2 记录每阶段候选、排名、分数、Boost、Rerank 与回退原因；TopK、RRF k、Boost上限和超时必须可配置。"],
    ], [850, 3700, 700, 4110])
    add_heading(doc, "6.5 服务主管与系统治理", 2)
    add_table(doc, ["编号", "需求", "优先级", "验收标准"], [
        ["FR-40", "服务主管查看运营总览、风险、SLA、工单和跨角色动态。", "P0", "高风险、即将超时、未分配事项可直接跳转到处理入口。"],
        ["FR-41", "系统管理员管理资源权限、审计链路和策略审批。", "P0", "记录敏感操作、权限变更和高风险拦截；可发现权限冲突与到期复核。"],
        ["FR-42", "安全要求包括脱敏、最小权限、敏感操作人工确认和可审计性。", "P0", "界面和接口均不得以 Agent 自动执行方式绕过身份核验或审批。"],
        ["FR-43", "退款、订单、物流、会员查询通过 D1 业务表提供真实结构化结果。", "P0", "工具不包含硬编码业务记录；订单不存在、权限不足、超时和数据缺失返回不同错误码；Agent 不得根据失败结果猜测。"],
        ["FR-44", "模型网关启动自检并在不可用时返回明确系统异常，不得静默生成伪客服回答。", "P0", "健康检查验证 Base URL、模型名、Key 和真实连通性；模型失败响应 mode=unavailable 并带 traceId。"],
        ["FR-45", "生产发布采用候选版本、迁移校验、灰度监控和可回滚策略。", "P0", "候选版本与当前版本可并存；关键链路异常时停止扩大流量；保留上一版镜像与数据库迁移。"],
        ["FR-46", "用户问候、闲聊和知识未命中问题同样调用真实 Agent 模型；没有可信知识时只做自然澄清，不编造业务事实。", "P0", "用户端正常回答返回 mode=model-grounded；模型不可用返回明确系统异常；不再以确定性文案伪装模型回答。"],
    ], [850, 3700, 700, 4110])

    add_heading(doc, "7. 数据、集成与非功能要求", 1)
    add_heading(doc, "7.1 核心数据对象", 2)
    add_table(doc, ["对象", "关键字段/关系", "治理要求"], [
        ["会话", "客户、渠道、意图、风险、SLA、历史、AI 回答、交接状态", "保留交接上下文；敏感字段脱敏；可关联工单与知识版本。"],
        ["知识", "编号、版本、内容、来源、责任团队、风险、有效期、状态、覆盖率", "来源可追溯；版本可回退；生效状态决定是否可被客服/Agent 引用。"],
        ["评测样本", "问题、标签、期望、模型输出、归因、复核结论", "保留原始样本与变更历史；不得由单人直接覆盖金标。"],
        ["策略", "路由规则、阈值、模型、风险动作、审批、版本", "生产变更必须审批、记录生效范围并支持回退。"],
        ["向量索引、影子运行与复核", "租户、Embedding 模型、当前索引版本、文档向量、traceId、候选分数、来源差异、耗时、失败码、复核结论与复核人", "新索引完整装载后才切换版本指针；复核按运行唯一且写审计；门禁只读，不得阻塞、替换回答或自动切流。"],
        ["审计事件", "主体、动作、资源、时间、结果、上下文", "覆盖权限、策略、知识和高风险业务操作。"],
    ], [1600, 4300, 3460])
    add_heading(doc, "7.2 外部依赖", 2)
    for item in [
        "订单、支付、物流、会员和工单系统提供可核验事实与业务动作状态。",
        "知识来源系统或文件库提供政策、SOP、FAQ 与来源快照。",
        "大模型、向量检索、意图识别和路由服务提供 Agent 能力。",
        "统一身份认证、权限系统、日志与监控平台提供企业治理能力。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "7.3 非功能与安全要求", 2)
    for item in [
        "可用性：生产核心会话与知识查询应具备监控、告警和降级方案。",
        "可追溯：每次回答应记录使用的知识、策略和模型版本。",
        "权限：按空间、角色和资源最小授权；角色切换不自动扩大权限。",
        "隐私：PII 默认脱敏；敏感操作需要身份核验与人工确认。",
        "可回退：知识和策略均需要保留线上版本、变更记录和回退路径。",
        "影子隔离：影子检索使用 Worker 后台任务保活；超时、索引缺失或写入失败只进入独立记录，不得延长或中断用户主链路。",
        "证据完整性：展示筛选与统计 cohort 分离；灰度门禁必须基于单一索引版本和四类业务覆盖，结果只用于决策提示。",
        "有界执行：Agent调用工具和模型前必须获得运行预算；不得自由循环，超时或预算耗尽不得继续使用迟到结果。",
        "Retrieval V2：检索范围、租户/状态/有效期/版本过滤必须在候选生成前生效；BM25 与 Dense 先独立宽召回，再使用 RRF、与RRF同量级的受限业务 Boost 和有超时的语义 Reranker；缺失向量不得进入Dense排名；Reranker 或 Dense 失败不得阻塞用户回答。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "8. 成功指标与埋点", 1)
    add_para(doc, "指标分为结果、质量、效率和安全四层。当前页面展示的数值是原型样例，正式目标值应在试点基线采集后确定。")
    add_table(doc, ["层级", "指标", "拟衡量的问题", "目标/阈值"], [
        ["北极星", "可信闭环解决率", "用户问题在有依据、无越权的前提下被解决的比例", "[Assumption] 试点后定义基线与提升目标"],
        ["服务结果", "72 小时未重开率、满意度、SLA 达标率", "人工与 AI 共同交付是否真正解决问题", "不低于现有人工服务基线"],
        ["AI 质量", "可信引用率、知识命中率、低置信度率、误转人工率", "回答是否有依据、是否正确选择自动/人工处理", "高风险误放行必须为 0"],
        ["知识治理", "知识缺口数、冲突处理时长、复核及时率、发布后改善幅度", "知识团队是否有效减少真实会话问题", "按业务类目定义 SLA"],
        ["效率/成本", "平均响应时延、单会话成本、建议采纳率", "质量提升是否以可接受成本实现", "不以降低安全为代价"],
        ["安全", "PII 脱敏覆盖、高风险拦截、越权访问、审计完整率", "企业级治理是否有效", "高风险自动执行为 0"],
        ["检索灰度证据", "完成样本、类目覆盖、变化复核、失败率、P95、改善率、退化率", "影子方案是否具备进入灰度设计的证据", "completed≥200；四类各≥25；明确结论变化≥30；失败率≤1%；P95≤500ms；退化率≤5%；改善率≥60%"],
        ["Retrieval V2", "Recall@5/10、Hit@5、MRR、nDCG@5、BM25/Dense/RRF/Rerank分阶段延迟", "宽召回与精排是否减少漏召回，同时保持可接受时延", "真实回归集V1/Hybrid-before/Rerank-after Recall@5为97.87%/96.81%/84.04%；新增集为96.47%/96.47%/88.24%。V2未通过晋级门禁，只能显式实验"],
        ["Agent 泛化", "新数据完整Agent通过率", "优化是否只记住开发评测题", "GLM-5.2控制、Step 9.7与Step 9.8均为83%；治理能力已完成但不宣称效果提升"],
    ], [1200, 2200, 3500, 2460])
    add_heading(doc, "8.1 建议埋点事件", 2)
    for item in [
        "会话：agent_answer_generated、agent_handoff_triggered、agent_answer_adopted、conversation_closed、conversation_reopened。",
        "知识：knowledge_gap_created、knowledge_test_passed、knowledge_review_submitted、knowledge_published、knowledge_rolled_back。",
        "评测/策略：evaluation_run_finished、failure_root_cause_marked、strategy_review_submitted、strategy_published。",
        "检索影子：retrieval_shadow_completed、retrieval_shadow_failed、retrieval_candidate_changed、retrieval_shadow_reviewed、embedding_index_activated。",
        "治理：permission_changed、sensitive_operation_blocked、audit_event_viewed。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "9. 发布、验收与风险", 1)
    add_heading(doc, "9.1 发布策略", 2)
    add_number(doc, "先在沙箱空间完成角色、权限、知识、策略和审计链路测试。")
    add_number(doc, "选择一个业务类目在生产空间灰度，限定知识与策略生效流量。")
    add_number(doc, "监控可信引用、风险转人工、投诉、SLA、人工反馈与成本。")
    add_number(doc, "达到指标和双人审批要求后扩大流量；发现高风险异常立即回退至上一版本。")
    add_number(doc, "影子证据门禁达标仅允许进入灰度设计；正式切流仍需冻结索引版本、审批、限流监控和回退演练。")
    add_heading(doc, "9.2 主要风险与缓解", 2)
    add_table(doc, ["风险", "影响", "缓解方案"], [
        ["知识口径错误或过期", "错误回答、投诉、合规风险", "来源、责任人、有效期、冲突检测、测试门禁、双人复核和回退。"],
        ["模型幻觉或越权承诺", "资金、赔付、账户等高风险事故", "高风险规则强制转人工；回答附依据；策略审批与审计。"],
        ["角色权限设计过宽", "敏感数据泄露或非授权发布", "空间隔离、最小权限、权限到期复核、审计追踪。"],
        ["只优化模型分数而忽略业务结果", "用户体验与成本恶化", "使用可信闭环解决率及服务/安全护栏指标共同评估。"],
        ["演示交互被误认为生产流程", "错误预期和治理缺失", "在上线计划中明确后端持久化、审批、通知和回退依赖。"],
        ["影子向量服务拖慢或干扰现网", "用户等待增加或稳定回答来源被意外替换", "后台 waitUntil 隔离、5 秒向量超时、异常吞吐隔离、开关回退和独立 trace 查询；生产回答继续使用稳定检索。"],
        ["筛选口径或少量样本造成虚假改善", "错误判断 Hybrid 已可灰度", "分析 cohort 与明细筛选分离；要求单一索引版本、四类覆盖、最少样本、人工复核和可解释阻塞项；门禁不自动切流。"],
        ["反复针对固定评测集优化", "已知集分数上升但真实新问法退化", "已知集只用于防回归；实现前冻结全局对话去重的新数据并做同集A/B；候选先过未见集，再运行旧集保护，任一门禁失败即停止；数据版本、指纹和失败归因必须留档，禁止为单题添加无业务依据的特例。"],
        ["Hybrid候选改变已验证行为", "新数据检索提升，但旧场景来源或转人工可能波动", "新数据完整Agent作为技术优化主指标；旧集用于发现明显退化并保持高风险100%；保留稳定来源回退。Step 9.1以83%成为技术基线，但普通生产仍不自动切换。"],
        ["静态证据门控过严", "有效语义候选被回退，完整Agent反而退化", "证据不足不得直接永久回退；只允许一次受约束查询改写和重检索，仍失败则回退基线。Step 9.2因83%降至79%被拒绝。"],
        ["纠正检索触发稀疏或无增益", "增加复杂度但最终回答不变", "纠正最多一次且保留Step 9.1首次结果；只承认原始问题证据。Step 9.3仅3/100触发、0条采用，完整Agent仍83%，因此不替换基线。"],
        ["多知识组合缺少真实采用机会", "增加证据拼接复杂度但不改变回答", "仅在主知识缺少用户明确需求槽位时组合；同租户/索引/类目/来源并检查重复与条件冲突。Step 9.4有94条进入判断但0条组合，完整Agent仍83%，不替换基线。"],
        ["工具或模型自由循环", "重复调用、成本失控、延迟扩大或使用超时结果", "有界状态机限制总步骤7、工具1次、模型1次和25秒；预算耗尽立即安全降级并记录原因。"],
        ["模型供应商参数不兼容", "输出异常、无效调用或不可比较评测", "模型切换建立独立运行版本；GLM-5.2使用非思考模式和供应商默认采样参数，健康检查与真实usage验收后再评测。"],
        ["宽召回或Reranker改变来源", "候选增加后错误知识进入回答，或远程Reranker超时增加延迟", "V1/普通生产隔离；V2保留全阶段审计；Reranker超时严格回退RRF+Boost；Dense失败回退可信BM25-only。真实评测中Reranker 0 fallback但使两套Recall@5分别下降12.77和8.23个百分点，因此拒绝晋级。"],
    ], [2200, 2700, 4460])

    add_heading(doc, "10. 开放问题与后续路线", 1)
    for item in [
        "[Open Question] 业务规则负责人、合规审批人的独立工作台、审批时限和替代人机制如何设计？",
        "[Open Question] 当前知识测试的真实样本集、金标维护机制、通过阈值和盲测规则如何确定？",
        "[Open Question] 多租户隔离、数据保留期限、导出权限与审计保留策略如何满足目标行业合规要求？",
        "[Open Question] Agent API 的模型供应商、故障降级、成本预算与观测标准如何确定？",
        "下一优先级：完成 AI 运营工作台的“失败归因 - 任务回流 - 策略评测 - 审批发布”真实闭环。",
        "回答能力技术路线已完成至有界Agentic RAG：GLM-5.2独立控制、统一证据包和有界状态机完整Agent均为83%。Phase 1.1真实Dense/RRF/Reranker已在回归集与独立新增集完成，但Hybrid未超过V1且Reranker显著退化，因此V2继续隔离、Phase 2不开始。",
        "后续迭代：将知识运营的原型状态接入后端持久化审批、真实灰度流量控制和可审计回退。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "11. PRD 维护与变更记录", 1)
    add_para(doc, "本文件是产品、设计、研发、数据、业务和合规协作的单一需求基线。每次修改系统时，产品负责人应先判断是否影响角色、流程、权限、指标、范围、依赖或风险；如影响，则在同一次变更中更新本 PRD 的相应章节、版本和变更记录。")
    add_table(doc, ["版本", "日期", "变更摘要", "状态"], [
        ["V1.0", "2026-08-14", "基于当前 ServicePilot 原型建立完整产品基线；纳入角色工作台和知识运营闭环。", "已创建"],
        ["V1.1", "2026-08-17", "将体验用户提问、Agent 处理、转人工队列和客服接管/回复/结案接入持久化会话与消息。", "已更新"],
        ["V1.2", "2026-08-17", "补充消费者端多角色、长期会话列表、完整历史恢复和客服消息回传；用户端隐藏 Agent 内部实现术语。", "已更新"],
        ["V1.3", "2026-08-17", "新增知识运营“话术训练”工作台：预设回复资产、自然表达预览、权限分工、双人复核和 Agent 训练影响说明。", "已更新"],
        ["V1.4", "2026-08-17", "将生产候选链路改为真实 Qwen、D1 业务表、租户隔离、来源/trace 追踪、非静默模型失败和自动工单；测试种子限定在测试租户。", "已更新"],
        ["V1.5", "2026-08-18", "问候、闲聊和知识未命中场景改为真实调用 Qwen；模型不可用时返回明确异常；同步验证用户端真实对话链路。", "已更新"],
        ["V1.6", "2026-08-18", "修复客服工作台与用户体验端租户不一致导致会话列表为空的问题；客服会话、工单、接管和回复统一读取测试租户。", "已更新"],
        ["V1.7", "2026-08-18", "新增测试租户知识资料批量导入链路：500 份带来源、版本、有效期、风险和合成标记的知识资产以待审核状态入库，生产租户保持空数据。", "已更新"],
        ["V1.8", "2026-08-18", "修复知识运营工作台待复核筛选无响应的问题；知识资产列表按测试租户真实状态查询并支持待复核、生效和草稿筛选。", "已更新"],
        ["V1.9", "2026-08-19", "完成全系统可操作性审计：补齐待办、通知、趋势筛选、审计搜索/风险筛选/CSV 导出、权限矩阵、资源血缘与菜单、评测优化任务等操作反馈；优化用户体验端双栏布局、响应式折叠、滚动边界和可读字号。", "已更新"],
        ["V2.0", "2026-08-21", "新增生产影子 Hybrid 检索可观测性：D1 版本化向量与当前版本指针、真实请求后台影子运行、候选分差/理由/延迟/失败记录及租户隔离的只读查询；不切换回答来源、Prompt 或转人工。", "已更新"],
        ["V2.1", "2026-08-21", "新增影子候选变化人工复核、复核权限与审计、按业务类目/索引版本统计，以及只读灰度证据门禁；门禁达标仅允许进入灰度设计，不自动切换生产流量。", "已更新"],
        ["V2.2", "2026-08-24", "Step 8.7动作门控在既有100条完整Agent开发回归集连续两轮达到94%；明确该成绩只作为已知集防回归基线，不能代表未知数据泛化。", "已更新"],
        ["V2.3", "2026-08-24", "生成模型切换至qwen3.7-plus-2026-05-26；新增100条CSDS新数据完整Agent泛化套件，冻结版本与指纹，首轮通过率81%，并将已知回归与泛化指标分开治理。", "已更新"],
        ["V2.4", "2026-08-24", "新增实现前冻结的第二批100条留出集和同集A/B门禁；Step 8.9留出集仅提升1个百分点但既有回归从94%降至88%，因此拒绝方案并恢复Step 8.7默认。", "已更新"],
        ["V2.5", "2026-08-24", "第三批100条全局对话去重留出集验证知识条目级动作风险矩阵：Step 8.7同集两次79%，Step 8.10为71%；第一道泛化门禁失败后停止旧集回归并拒绝候选。", "已更新"],
        ["V2.6", "2026-08-24", "将版本化受约束Hybrid接入显式完整Agent候选分支：新数据79%提升至83%，旧回归94%变为93%，高风险保持100%。后续按新数据主指标口径将83%确立为技术基线，普通生产仍未切换。", "已更新"],
        ["V2.7", "2026-08-24", "新增静态检索证据充分性门控；完整Agent从83%降至79%，因此拒绝Step 9.2并保留83%基线。下一步仅允许一次受约束Corrective RAG，禁止自由重试。", "已更新"],
        ["V2.8", "2026-08-24", "完成一次有界Corrective RAG：3/100触发、0条采用，完整Agent保持83%且无提升，因此不替换Step 9.1基线；下一步进入受约束多知识组合。", "已更新"],
        ["V2.9", "2026-08-24", "完成受约束多知识组合：94/100进入判断、0条实际组合，完整Agent保持83%且无提升；补齐多来源消息/运行审计，下一步进入工具决策与回答审查。", "已更新"],
        ["V3.0", "2026-08-25", "生成模型切换GLM-5.2非思考模式；完成确定性工具决策、声明审查、知识/工具统一证据包和有界Agentic RAG。GLM控制、Step 9.7与9.8均为83%，普通生产仍未启用候选。", "已更新"],
        ["V3.1", "2026-08-25", "执行 Step 9.8 Phase 1 Retrieval Foundation：新增 RetrievalScope、BM25/Dense 宽召回、RRF、受限业务 Boost、供应商无关语义 Reranker 与分阶段 Trace。V2显式隔离；本地BM25-only Recall@5为77.66%，低于V1的95.74%，Dense/Reranker未运行，因此不切流、不宣称提升。", "已更新"],
        ["V3.2", "2026-08-25", "完成Phase 1.1真实模型检索评测：live D1知识与当前向量索引双向一致，真实Embedding和qwen3-rerank均0 fallback。两套数据中Hybrid-before未超过V1，Reranker使Recall@5明显下降，因此V2拒绝晋级、继续隔离，Phase 2不开始。", "已更新"],
    ], [1200, 1500, 5160, 1500])
    add_heading(doc, "12. 参考模板与来源", 1)
    add_para(doc, "PRD 结构参考成熟企业协作模板的通用做法：在同一文档中维护目标、成功指标、假设、用户故事、范围、开放问题和支持材料，并把 PRD 作为持续更新的协作基线。")
    for source in [
        "Atlassian Confluence, Product requirements document (PRD) template: https://www.atlassian.com/software/confluence/templates/product-requirements",
        "Atlassian Agile Product Management, How to create a PRD: https://www.atlassian.com/agile/product-management/requirements/",
        "ServicePilot 当前代码与已部署原型（角色工作台、会话、工单、知识运营、质检、策略、权限与审计）。",
    ]:
        add_bullet(doc, source)

    doc.core_properties.title = "ServicePilot 产品需求文档（PRD）"
    doc.core_properties.subject = "企业智能服务运营平台产品基线"
    doc.core_properties.author = "ServicePilot 产品团队"
    doc.core_properties.comments = "Living PRD - update with every product change"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)

if __name__ == "__main__":
    main()
